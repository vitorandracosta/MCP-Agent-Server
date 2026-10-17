"""
Servidor MCP "escrita": cria arquivos novos, SOMENTE dentro de
config.OUTPUT_DIR (pasta saidas/). Não sobrescreve nada sem pedido
explícito e não aceita caminhos que escapem da pasta.
"""
from pathlib import Path

from mcp.server.fastmcp import FastMCP

import config
from servers.seguranca import dentro_de

mcp = FastMCP("escrita")

EXTENSOES_PERMITIDAS = {".md", ".txt", ".json", ".csv", ".py", ".html", ".docx"}
MAX_CHARS = 200_000


# Erros são levantados como exceção: o FastMCP os devolve com isError=True,
# e o orquestrador consegue distinguir falha de sucesso.
def _destino(nome: str) -> Path:
    nome = nome.strip().strip('"').strip("'").replace("\\", "/")
    if not Path(nome).suffix:
        nome += ".md"
    alvo = (config.OUTPUT_DIR / nome).resolve()
    if not dentro_de(alvo, [config.OUTPUT_DIR]) or alvo == config.OUTPUT_DIR:
        raise ValueError("Nome inválido: o arquivo precisa ficar dentro da pasta saidas/.")
    if alvo.suffix.lower() not in EXTENSOES_PERMITIDAS:
        raise ValueError(f"Extensão não permitida. Use uma de: {', '.join(sorted(EXTENSOES_PERMITIDAS))}.")
    return alvo


@mcp.tool()
def criar_arquivo(nome: str, conteudo: str, sobrescrever: bool = False) -> str:
    """Cria um arquivo novo na pasta saidas/ (ex.: nome='resumo_calculo.md'). Aceita .md, .txt, .json, .csv, .py, .html e .docx; sem extensão, usa .md."""
    alvo = _destino(nome)
    if alvo.exists() and not sobrescrever:
        raise ValueError(f"Já existe {alvo.name}. Escolha outro nome ou use sobrescrever=true.")
    if len(conteudo) > MAX_CHARS:
        raise ValueError(f"Conteúdo grande demais ({len(conteudo)} caracteres; limite {MAX_CHARS}).")

    alvo.parent.mkdir(parents=True, exist_ok=True)
    if alvo.suffix.lower() == ".docx":
        import docx
        documento = docx.Document()
        for linha in conteudo.splitlines():
            if linha.startswith("# "):
                documento.add_heading(linha[2:], level=1)
            elif linha.startswith("## "):
                documento.add_heading(linha[3:], level=2)
            else:
                documento.add_paragraph(linha)
        documento.save(str(alvo))
    else:
        alvo.write_text(conteudo, encoding="utf-8")
    return f"Arquivo criado: {alvo} ({alvo.stat().st_size} bytes)"


def _arquivos_gerados() -> str:
    if not config.OUTPUT_DIR.exists():
        return "(nenhum arquivo gerado ainda)"
    arquivos = [p for p in sorted(config.OUTPUT_DIR.rglob("*"))
                if p.is_file() and not p.name.startswith(".")]
    if not arquivos:
        return "(nenhum arquivo gerado ainda)"
    return "\n".join(f"{p.relative_to(config.OUTPUT_DIR)} ({p.stat().st_size} bytes)" for p in arquivos)


@mcp.resource("saidas://lista", name="Arquivos gerados", mime_type="text/plain")
def resource_arquivos_gerados() -> str:
    """Arquivos já existentes em saidas/ (o escritor usa para não repetir nomes)."""
    return _arquivos_gerados()


@mcp.prompt(name="relatorio", description="Pede um relatório sobre um assunto, salvo como arquivo.")
def prompt_relatorio(assunto: str, formato: str = "md") -> str:
    return (f"Pesquise nos meus documentos locais sobre '{assunto}' e crie um relatório "
            f"em formato .{formato} com introdução, tópicos principais e lista de fontes. "
            "Salve com criar_arquivo.")


@mcp.tool()
def listar_arquivos_gerados() -> str:
    """Lista os arquivos já criados na pasta saidas/."""
    return _arquivos_gerados()


if __name__ == "__main__":
    mcp.run(transport="stdio")
